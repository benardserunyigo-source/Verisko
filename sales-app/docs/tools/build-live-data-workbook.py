from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule

KEY = ""  # leave blank; the Owner pastes the export key into Setup!C4 after upload
APP = "https://verisko-sales-2026.netlify.app"
NAVY = "012169"; TEAL = "0B7285"; LIGHT = "EAF0FA"; GREY = "F4F5F7"; MID = "D9DEE8"; WHITE = "FFFFFF"; INK = "1F2937"; MUTED = "6B7280"; INPUT = "FFF4CC"
F = "Arial"
def font(**k): k.setdefault("name", F); return Font(**k)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=MID)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------------------------------------------------------------- helpers
def col(tab, header, bounded=True):
    rng = "$A$2:$ZZ$10000" if bounded else "$A:$ZZ"
    return 'INDEX(%s!%s,0,MATCH("%s",%s!$1:$1,0))' % (tab, rng, header, tab)
def countif(tab, header, crit): return '=IFERROR(COUNTIF(%s,%s),0)' % (col(tab, header), crit)
def total(tab): return '=IFERROR(MAX(0,COUNTA(%s!$A$1:$A$10000)-1),0)' % tab
def header_row(ws, row, cells, widths=None):
    for i, text in enumerate(cells, start=1):
        c = ws.cell(row, i, text); c.font = font(bold=True, color=WHITE, size=10); c.fill = fill(NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True); c.border = box

# ---------------------------------------------------------------- Dashboard
d = wb.active; d.title = "Dashboard"
d.sheet_view.showGridLines = False
for c in range(1, 16): d.column_dimensions[get_column_letter(c)].width = 11.5
d.column_dimensions["A"].width = 3
d.merge_cells("B1:N1"); d["B1"] = "Verisko Uganda Operations · Live dashboard"
d["B1"].font = font(bold=True, size=20, color=NAVY); d.row_dimensions[1].height = 34
d.merge_cells("B2:N2")
d["B2"] = '="Live from the Verisko app · refreshes about hourly and on open · opened "&TEXT(NOW(),"d mmm yyyy, hh:mm")'
d["B2"].font = font(size=10, color=MUTED, italic=True)

tiles = [
    ("Prospects", "='Chart data'!$B$12", "all time"),
    ("Awaiting review", "='Chart data'!$U$3", "Operations to check"),
    ("Approved prospects", "='Chart data'!$U$4", "verified by Operations"),
    ("Closed sales", "=IFERROR(COUNTIF(%s,\"yes\"),0)" % col("Prospects", "closed_sale"), "commission paid on these"),
    ("Confirmed site visits", "='Chart data'!$K$4", "ready for handoff"),
    ("Float balance (UGX)", "='Chart data'!$N$10-'Chart data'!$O$10", "approved in minus approved out"),
    ("Cash awaiting approval", "=IFERROR(COUNTIF(%s,\"pending\"),0)" % col("'Cash flow'", "status"), "entries for the Owner"),
    ("Open jobs", "='Chart data'!$R$13", "not yet handed over"),
]
starts = ["B", "E", "H", "K"]
for i, (label, formula, sub) in enumerate(tiles):
    r = 4 if i < 4 else 8
    s = starts[i % 4]; e = chr(ord(s) + 2)
    for rr in (r, r + 1, r + 2):
        for cc in range(ord(s) - 64, ord(e) - 64 + 1):
            cell = d.cell(rr, cc); cell.fill = fill(LIGHT if i % 2 == 0 else GREY)
    d.merge_cells(f"{s}{r}:{e}{r}"); d.merge_cells(f"{s}{r+1}:{e}{r+1}"); d.merge_cells(f"{s}{r+2}:{e}{r+2}")
    d[f"{s}{r}"] = label; d[f"{s}{r}"].font = font(bold=True, size=10, color=MUTED); d[f"{s}{r}"].alignment = Alignment(horizontal="center", vertical="bottom")
    d[f"{s}{r+1}"] = formula; d[f"{s}{r+1}"].font = font(bold=True, size=24, color=NAVY); d[f"{s}{r+1}"].alignment = Alignment(horizontal="center", vertical="center"); d[f"{s}{r+1}"].number_format = "#,##0"
    d[f"{s}{r+2}"] = sub; d[f"{s}{r+2}"].font = font(size=9, color=MUTED); d[f"{s}{r+2}"].alignment = Alignment(horizontal="center", vertical="top")
    d.row_dimensions[r + 1].height = 36
for r in (7, 11): d.row_dimensions[r].height = 8

def section(ws, cell, text):
    ws[cell] = text; ws[cell].font = font(bold=True, size=12, color=NAVY)
section(d, "B13", "Sales pipeline by stage"); section(d, "I13", "Prospects by salesperson")
section(d, "B31", "Prospects by business type"); section(d, "I31", "Site visits by status")
section(d, "B49", "Cash flow, last 6 months (UGX)"); section(d, "I49", "Jobs by stage")
d["B67"] = "How to read this"; d["B67"].font = font(bold=True, size=12, color=NAVY)
notes = [
    "Every number is a formula over the data tabs, which pull straight from the app. Nothing is typed by hand.",
    "Float balance = approved money in − approved money out. Entries awaiting the Owner's approval are not in it yet.",
    "Chart figures live on the 'Chart data' tab. Records are edited in the app, not here.",
]
for i, n in enumerate(notes, start=68):
    d.merge_cells(f"B{i}:N{i}"); d[f"B{i}"] = "• " + n; d[f"B{i}"].font = font(size=10, color=INK); d[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top"); d.row_dimensions[i].height = 28

# ---------------------------------------------------------------- Setup
s = wb.create_sheet("Setup")
s.sheet_view.showGridLines = False
s.column_dimensions["A"].width = 3; s.column_dimensions["B"].width = 36; s.column_dimensions["C"].width = 100
s["B1"] = "Setup"; s["B1"].font = font(bold=True, size=16, color=NAVY)
s["B2"] = "Only the Owner needs this tab. Team members can use the Dashboard and data tabs as they are."; s["B2"].font = font(size=10, color=MUTED, italic=True)
s["B4"] = "Export key"; s["C4"] = KEY; s["C4"].fill = fill(INPUT); s["C4"].font = font(color="0000FF"); s["C4"].border = box
s["B5"] = "App address"; s["C5"] = APP; s["C5"].fill = fill(INPUT); s["C5"].font = font(color="0000FF"); s["C5"].border = box
s["B6"] = "Feed address (works itself out)"; s["C6"] = '=IF(C4="","Paste the export key in C4",C5&"/api/export?key="&C4)'; s["C6"].font = font(color=MUTED, size=9)
for r in (4, 5, 6): s[f"B{r}"].font = font(bold=True, size=10, color=INK)
s["B8"] = "How it works"; s["B8"].font = font(bold=True, size=12, color=NAVY)
steps = [
    "Key: app → Settings → Live data export → Generate export key → Copy. Paste it in C4.",
    "First open: click 'Allow access' when Google asks. Tabs then fill and refresh themselves.",
    "Refresh: about hourly and on every open. For an instant refresh, edit C4 (add a space, remove it).",
    "Read-only: change records in the app, not here. Photo links open the stored image.",
    "Privacy: the key reads all company data. Share this file as Viewer only. Leaked key? App → Settings → Generate new key, paste it here.",
    "Snapshot: File → Download → Microsoft Excel.",
]
for i, t in enumerate(steps, start=9):
    s[f"B{i}"] = f"{i-8}."; s[f"B{i}"].alignment = Alignment(horizontal="right", vertical="top"); s[f"B{i}"].font = font(bold=True, color=NAVY)
    s[f"C{i}"] = t; s[f"C{i}"].alignment = Alignment(wrap_text=True, vertical="top"); s[f"C{i}"].font = font(size=10, color=INK); s.row_dimensions[i].height = 30
s["B16"] = "Tabs"; s["B16"].font = font(bold=True, size=12, color=NAVY)
tabs_desc = [("Dashboard", "Headline numbers and charts."), ("Prospects", "Every lead, with GPS pin, photo and review status."), ("Visits", "Site visits booked for Operations."), ("Follow-ups", "Calls and visits logged per prospect."), ("Jobs", "Quotes and installations, Draft to Handed over."), ("Cash flow", "Every cash entry with status and receipt link."), ("Team", "Access and roles."), ("Technicians", "Installer roster."), ("Chart data", "Tables behind the charts.")]
for i, (t, desc) in enumerate(tabs_desc, start=17):
    s[f"B{i}"] = t; s[f"B{i}"].font = font(bold=True, size=10, color=INK); s[f"C{i}"] = desc; s[f"C{i}"].font = font(size=10, color=INK)

# ---------------------------------------------------------------- Data tabs
# (tab title, feed table, [(header, width, number_format)])
W = {"id": 16, "business": 28, "type": 14, "contact": 18, "phone": 17, "location": 22, "source": 16, "spoke_to_decision_maker": 12, "existing_cameras": 10, "budget": 14, "stage": 20, "next_action": 28, "follow_up_date": 13, "security_concern": 30, "areas_to_cover": 22, "notes": 30, "created": 12, "created_by": 14, "created_by_email": 24, "review_status": 13, "reviewed_by": 14, "reviewed_at": 12, "review_note": 26, "closed_sale": 10, "closed_by": 14, "closed_at": 12, "gps_lat": 11, "gps_lng": 11, "gps_accuracy_m": 10, "gps_captured_at": 20, "gps_map_link": 34, "photo_link": 34, "follow_ups_count": 10, "last_follow_up_at": 20, "last_follow_up_note": 30,
     "prospect_id": 16, "date": 12, "time": 8, "operations_owner": 16, "status": 13, "purpose": 20, "directions": 34, "n": 5, "at": 20, "by": 14, "by_email": 24, "note": 34,
     "ref": 12, "created_at": 12, "final_price": 14, "materials_count": 9, "materials_total": 14, "materials": 40,
     "direction": 10, "amount": 14, "category": 18, "method": 12, "job_id": 16, "preapproved": 11, "receipt_link": 34,
     "name": 20, "email": 26, "role": 12, "skills": 24, "active": 8}
NUM = {"amount": "#,##0", "final_price": "#,##0", "materials_total": "#,##0", "budget": "@", "gps_lat": "0.000000", "gps_lng": "0.000000", "phone": "@", "id": "@", "prospect_id": "@", "job_id": "@"}
COLS = {
 "prospects": ["id","business","type","contact","phone","location","source","spoke_to_decision_maker","existing_cameras","budget","stage","next_action","follow_up_date","security_concern","areas_to_cover","notes","created","created_by","created_by_email","review_status","reviewed_by","reviewed_at","review_note","closed_sale","closed_by","closed_at","gps_lat","gps_lng","gps_accuracy_m","gps_captured_at","gps_map_link","photo_link","follow_ups_count","last_follow_up_at","last_follow_up_note"],
 "visits": ["id","prospect_id","business","contact","phone","location","date","time","operations_owner","status","purpose","directions"],
 "followups": ["prospect_id","business","n","at","by","by_email","note","gps_lat","gps_lng","gps_accuracy_m","gps_captured_at","gps_map_link"],
 "jobs": ["id","ref","stage","prospect_id","business","created_at","created_by","created_by_email","final_price","materials_count","materials_total","materials"],
 "transactions": ["id","date","direction","amount","category","method","prospect_id","business","job_id","note","preapproved","status","created_by","created_by_email","created_at","reviewed_by","reviewed_at","review_note","receipt_link"],
 "users": ["id","name","email","role","created"],
 "technicians": ["id","name","phone","skills","active","created_at"],
}
DATA_TABS = [("Prospects","prospects"),("Visits","visits"),("Follow-ups","followups"),("Jobs","jobs"),("Cash flow","transactions"),("Team","users"),("Technicians","technicians")]
for title, table in DATA_TABS:
    t = wb.create_sheet(title)
    cols = COLS[table]; n = len(cols) + 6
    t["A1"] = '=IFERROR(IMPORTDATA(Setup!$C$6&"&table=%s"),"Waiting for data: check the Setup tab and click Allow access.")' % table
    for i in range(1, n + 1):
        c = t.cell(1, i); c.font = font(bold=True, color=WHITE, size=10); c.fill = fill(NAVY); c.alignment = Alignment(vertical="center"); c.border = box
        t.column_dimensions[get_column_letter(i)].width = W.get(cols[i-1], 14) if i <= len(cols) else 14
        fmt = NUM.get(cols[i-1]) if i <= len(cols) else None
        if fmt and fmt != "@":
            t.column_dimensions[get_column_letter(i)].number_format = fmt
    t.row_dimensions[1].height = 22
    t.freeze_panes = "C2" if table in ("prospects", "visits", "jobs") else "B2"
    last = get_column_letter(n)
    t.conditional_formatting.add(f"A2:{last}1000", FormulaRule(formula=['AND($A2<>"",MOD(ROW(),2)=0)'], fill=fill(GREY)))
    t.sheet_properties.tabColor = TEAL if table in ("prospects","visits","followups") else NAVY

# ---------------------------------------------------------------- Chart data
c = wb.create_sheet("Chart data")
c.sheet_view.showGridLines = False
c["A1"] = "Chart data · the small tables behind the Dashboard. All formulas over the data tabs."; c["A1"].font = font(bold=True, size=12, color=NAVY)
STAGES = ["New prospect", "Contact attempted", "Qualified", "Appointment proposed", "Appointment confirmed", "Lost", "Postponed"]
VERTICALS = ["Pharmacy", "Clinic", "Hospital", "Mobile money", "Retail shop", "Supermarket", "School", "Office", "Warehouse", "Residence", "Other"]
APPT = ["Proposed", "Confirmed", "Completed", "Rescheduled", "Cancelled", "No-show"]
JOBS = ["Draft", "Sent", "Accepted", "Scheduled", "In progress", "Installed", "Handed over", "Rejected", "Cancelled"]
def block(ws, col0, hdrs, labels, formulas, widths):
    for j, h in enumerate(hdrs):
        cell = ws.cell(2, col0 + j, h); cell.font = font(bold=True, color=WHITE, size=10); cell.fill = fill(NAVY); cell.border = box
        ws.column_dimensions[get_column_letter(col0 + j)].width = widths[j]
    for i, lab in enumerate(labels):
        r = 3 + i
        ws.cell(r, col0, lab).border = box
        for j, f in enumerate(formulas):
            cell = ws.cell(r, col0 + 1 + j, f(r, lab)); cell.number_format = "#,##0"; cell.border = box
    return 3 + len(labels)
# A:B pipeline
end = block(c, 1, ["Stage", "Prospects"], STAGES, [lambda r, lab: countif("Prospects", "stage", f"$A{r}")], [22, 11])
c.cell(end, 1, "Total").font = font(bold=True); c.cell(end, 2, f"=SUM(B3:B{end-1})").font = font(bold=True); c.cell(end, 2).number_format = "#,##0"
c.cell(11, 1, "Prospects (all)").font = font(bold=True); c.cell(12, 1, "Total"); c.cell(12, 2, total("Prospects")).number_format = "#,##0"
# D:E salesperson
for j, h in enumerate(["Salesperson", "Prospects"]):
    cell = c.cell(2, 4 + j, h); cell.font = font(bold=True, color=WHITE, size=10); cell.fill = fill(NAVY); cell.border = box
c.column_dimensions["D"].width = 20; c.column_dimensions["E"].width = 11
cb = col("Prospects", "created_by")
# One row per team member (from the Team tab), so the chart series is always numeric.
tn = col("Team", "name")
for r in range(3, 15):
    c.cell(r, 4, '=IFERROR(INDEX(%s,%d),"")' % (tn, r - 2)).border = box
    e = c.cell(r, 5, '=IF(D%d="",0,IFERROR(COUNTIF(%s,D%d),0))' % (r, cb, r)); e.border = box; e.number_format = "#,##0"
# G:H business type
block(c, 7, ["Business type", "Prospects"], VERTICALS, [lambda r, lab: countif("Prospects", "type", f"$G{r}")], [16, 11])
# J:K visit status
block(c, 10, ["Visit status", "Visits"], APPT, [lambda r, lab: countif("Visits", "status", f"$J{r}")], [14, 9])
# M:O cash by month
for j, h in enumerate(["Month", "Money in", "Money out"]):
    cell = c.cell(2, 13 + j, h); cell.font = font(bold=True, color=WHITE, size=10); cell.fill = fill(NAVY); cell.border = box
c.column_dimensions["M"].width = 10; c.column_dimensions["N"].width = 14; c.column_dimensions["O"].width = 14
dcol = col("'Cash flow'", "date"); dirc = col("'Cash flow'", "direction"); amt = col("'Cash flow'", "amount"); st = col("'Cash flow'", "status")
for i in range(6):
    r = 3 + i
    c.cell(r, 13, '=TEXT(EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),%d),"yyyy-mm")' % (i - 5)).border = box
    c.cell(r, 14, '=IFERROR(SUMPRODUCT((TEXT(%s,"yyyy-mm")=$M%d)*(%s="in")*(%s="approved"),%s),0)' % (dcol, r, dirc, st, amt)).number_format = "#,##0"
    c.cell(r, 15, '=IFERROR(SUMPRODUCT((TEXT(%s,"yyyy-mm")=$M%d)*(%s="out")*(%s="approved"),%s),0)' % (dcol, r, dirc, st, amt)).number_format = "#,##0"
    c.cell(r, 14).border = box; c.cell(r, 15).border = box
c.cell(10, 13, "All time (approved)").font = font(bold=True)
c.cell(10, 14, '=IFERROR(SUMIFS(%s,%s,"in",%s,"approved"),0)' % (amt, dirc, st)).number_format = "#,##0"
c.cell(10, 15, '=IFERROR(SUMIFS(%s,%s,"out",%s,"approved"),0)' % (amt, dirc, st)).number_format = "#,##0"
# Q:R jobs by stage
end = block(c, 17, ["Job stage", "Jobs"], JOBS, [lambda r, lab: countif("Jobs", "stage", f"$Q{r}")], [16, 9])
c.cell(13, 17, "Open (not handed over / rejected / cancelled)").font = font(bold=True); c.column_dimensions["Q"].width = 16
c.cell(13, 18, "=SUM(R3:R8)").number_format = "#,##0"
# T:U review status
block(c, 20, ["Review status", "Prospects"], ["pending", "approved", "query"], [lambda r, lab: countif("Prospects", "review_status", f"$T{r}")], [14, 11])
c.cell(6, 20, "(query = sent back to Sales)").font = font(size=9, color=MUTED)

# ---------------------------------------------------------------- Charts on Dashboard
def bar(title, cats, vals, horizontal=False, colors=None):
    ch = BarChart(); ch.type = "bar" if horizontal else "col"; ch.title = title; ch.style = 10
    ch.add_data(vals, titles_from_data=True); ch.set_categories(cats)
    ch.legend = None; ch.y_axis.majorGridlines = None; ch.height = 8.2; ch.width = 12.2
    ch.series[0].graphicalProperties.solidFill = NAVY
    if horizontal: ch.x_axis.scaling.orientation = "maxMin"
    return ch
cd = wb["Chart data"]
ch = bar("Prospects by stage", Reference(cd, min_col=1, min_row=3, max_row=9), Reference(cd, min_col=2, min_row=2, max_row=9), horizontal=True); d.add_chart(ch, "B14")
ch = bar("Prospects by salesperson", Reference(cd, min_col=4, min_row=3, max_row=14), Reference(cd, min_col=5, min_row=2, max_row=14)); d.add_chart(ch, "I14")
pie = PieChart(); pie.title = "Prospects by business type"; pie.add_data(Reference(cd, min_col=8, min_row=2, max_row=13), titles_from_data=True); pie.set_categories(Reference(cd, min_col=7, min_row=3, max_row=13)); pie.height = 8.2; pie.width = 12.2; pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True; d.add_chart(pie, "B32")
for i, lab in enumerate(APPT):
    r = 33 + i
    d[f"I{r}"] = lab; d[f"I{r}"].font = font(size=10, color=INK); d.merge_cells(f"I{r}:K{r}")
    d[f"L{r}"] = f"='Chart data'!$K${3+i}"; d[f"L{r}"].number_format = "#,##0"; d[f"L{r}"].font = font(bold=True, size=10, color=NAVY); d[f"L{r}"].alignment = Alignment(horizontal="right")
    for cc in "IJKL": d[f"{cc}{r}"].fill = fill(LIGHT if i % 2 == 0 else WHITE); d[f"{cc}{r}"].border = Border(bottom=thin)
cash = BarChart(); cash.type = "col"; cash.title = "Approved cash in vs out, by month"; cash.style = 10; cash.height = 8.2; cash.width = 12.2
cash.add_data(Reference(cd, min_col=14, min_row=2, max_col=15, max_row=8), titles_from_data=True); cash.set_categories(Reference(cd, min_col=13, min_row=3, max_row=8))
cash.series[0].graphicalProperties.solidFill = TEAL; cash.series[1].graphicalProperties.solidFill = NAVY; cash.y_axis.majorGridlines = None
d.add_chart(cash, "B50")
for i, lab in enumerate(JOBS):
    r = 51 + i
    d[f"I{r}"] = lab; d[f"I{r}"].font = font(size=10, color=INK); d.merge_cells(f"I{r}:K{r}")
    d[f"L{r}"] = f"='Chart data'!$R${3+i}"; d[f"L{r}"].number_format = "#,##0"; d[f"L{r}"].font = font(bold=True, size=10, color=NAVY); d[f"L{r}"].alignment = Alignment(horizontal="right")
    for cc in "IJKL": d[f"{cc}{r}"].fill = fill(LIGHT if i % 2 == 0 else WHITE); d[f"{cc}{r}"].border = Border(bottom=thin)

wb["Dashboard"].sheet_properties.tabColor = NAVY; wb["Setup"].sheet_properties.tabColor = "F59E0B"; wb["Chart data"].sheet_properties.tabColor = "9CA3AF"
wb.save("Verisko Live Data v2.xlsx")
import os; print("bytes", os.path.getsize("Verisko Live Data v2.xlsx"))
